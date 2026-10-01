"""Fixed synthetic Office inputs; document content never comes from user files."""
from pathlib import Path
import hashlib, json, random, sys
from importlib.metadata import version
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from PIL import Image, ImageDraw

def deterministic_package(path):
 """Fix package metadata after writers that stamp the current clock during save."""
 source=BytesIO(path.read_bytes()); output=BytesIO()
 with ZipFile(source) as old, ZipFile(output,'w',compression=ZIP_DEFLATED,compresslevel=9) as new:
  for name in sorted(old.namelist()):
   data=old.read(name)
   if name=='docProps/core.xml':
    properties=ElementTree.fromstring(data)
    for key in ('created','modified'):
     field=properties.find('{http://purl.org/dc/terms/}'+key)
     if field is not None: field.text='2000-01-01T00:00:00Z'
    data=ElementTree.tostring(properties,encoding='utf-8',xml_declaration=True)
   entry=ZipInfo(name,date_time=(1980,1,1,0,0,0)); entry.compress_type=ZIP_DEFLATED
   entry.create_system=3; entry.external_attr=0o600<<16
   new.writestr(entry,data,compress_type=ZIP_DEFLATED,compresslevel=9)
 path.write_bytes(output.getvalue())

root=Path(sys.argv[1]).resolve(); root.mkdir(parents=True, exist_ok=True)
out=root/'inputs'; out.mkdir(exist_ok=True)
rng=random.Random(20260911)
img=Image.new('RGB',(1600,900)); px=img.load()
for y in range(900):
 for x in range(1600):
  n=rng.randrange(32); px[x,y]=((x//7+n)%256,(y//4+n)%256,(x//13+y//8+n)%256)
draw=ImageDraw.Draw(img)
for i in range(12): draw.rectangle((60+i*120,700-i*40,140+i*120,830),fill=(30,100+i*9,190))
img.save(out/'chart.png'); rows=[]
for count in (3,30):
 d=Document(); d.styles['Normal'].font.name='Arial'; d.styles['Normal'].font.size=Pt(11)
 for i in range(count):
  if i: d.add_page_break()
  d.add_heading(f'Office preview benchmark {i+1}',level=1)
  p=d.add_paragraph('中文排版测试：文档预览、表格与图片。 ' + 'Deterministic synthetic business report. '*8)
  for run in p.runs:
   rpr=run._element.get_or_add_rPr(); fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'),'Songti SC'); rpr.append(fonts)
  t=d.add_table(rows=1,cols=4); t.style='Table Grid'
  for j,c in enumerate(t.rows[0].cells): c.text=['Item','Quantity','Price','Total'][j]
  for j in range(8):
   for k,c in enumerate(t.add_row().cells): c.text=[f'Product {j}',str(j+1),'12.50',str((j+1)*12.5)][k]
  d.add_picture(str(out/'chart.png'),width=Inches(5))
 name=f'report-{count}p.docx'; d.save(out/name); rows.append({'file':name,'dimensions':f'{count} explicit pages, Chinese/English, tables, shared 1600x900 image'})
for count in (200,3000):
 w=Workbook(); s=w.active; s.title='Sales'; s.append(['Index','Region','Product','Units','Price','Total'])
 for i in range(1,count+1): s.append([i,['East','West','North','South'][i%4],f'Item {i%50}',i%25+1,12.5,f'=D{i+1}*E{i+1}'])
 for cell in s[1]: cell.font=Font(name='Arial',bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='285A90')
 for col in 'ABCDEF': s.column_dimensions[col].width=15
 s.sheet_properties.pageSetUpPr.fitToPage=True; s.page_setup.orientation='landscape'; s.page_setup.paperSize=s.PAPERSIZE_A4; s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0
 s.print_title_rows='1:1'; s.print_area=f'A1:F{count+1}'
 name=f'sales-{count}r.xlsx'; w.save(out/name); rows.append({'file':name,'dimensions':f'{count} rows x 6 columns, formulas, fit width 1 page'})
for count in (5,30):
 p=Presentation(); p.slide_width=PInches(13.333); p.slide_height=PInches(7.5)
 for i in range(count):
  slide=p.slides.add_slide(p.slide_layouts[6]); box=slide.shapes.add_textbox(PInches(.6),PInches(.3),PInches(12),PInches(.8)); run=box.text_frame.paragraphs[0].add_run(); run.text=f'Slide {i+1}: Quarterly report 中文预览'; run.font.name='Arial'; run.font.size=PPt(28)
  slide.shapes.add_picture(str(out/'chart.png'),PInches(.7),PInches(1.5),width=PInches(8))
  tf=slide.shapes.add_textbox(PInches(9),PInches(2),PInches(3.5),PInches(4)).text_frame
  for j in range(6): tf.add_paragraph().text=f'Metric {j+1}: {100+i*7+j}'
 name=f'deck-{count}s.pptx'; p.save(out/name); rows.append({'file':name,'dimensions':f'{count} slides, shared 1600x900 image, text'})
for row in rows:
 deterministic_package(out/row['file'])
 row['bytes']=(out/row['file']).stat().st_size
 row['sha256']=hashlib.sha256((out/row['file']).read_bytes()).hexdigest()
(root/'fixtures.json').write_text(json.dumps(rows,indent=2)+'\n',encoding='utf-8')
(root/'generator.json').write_text(json.dumps({'python':sys.version,'packages':{name:version(name) for name in ['python-docx','openpyxl','python-pptx','pillow']},'seed':20260911,'zipTimestamp':[1980,1,1,0,0,0],'documentTimestamp':'2000-01-01T00:00:00Z'},indent=2)+'\n',encoding='utf-8')
print(json.dumps(rows,indent=2))
